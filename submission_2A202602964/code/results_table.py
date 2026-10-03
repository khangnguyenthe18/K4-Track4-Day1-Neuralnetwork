"""results_table.py — Lưu kết quả JSON và xuất file experiments.xlsx theo mẫu.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx.

Tên cột của sheet "Experiments":
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính)
"""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl

EXPERIMENT_COLUMNS = [
    "exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay",
    "batch", "epochs", "hidden", "dropout", "clip_norm", "precision", "init",
    "seed", "step0_loss", "best_val_loss", "best_epoch", "final_train_loss",
    "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB",
    "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes"
]


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra JSON."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    out_file = Path(results_dir) / f"{exp_id}.json"

    data_to_save = {
        "cfg": result.get("cfg", {}),
        "history": result.get("history", {}),
        "summary": result.get("summary", {}),
    }

    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2, ensure_ascii=False)

    return str(out_file)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    p = Path(results_dir)
    if not p.exists():
        return []
    results = []
    for f in sorted(p.glob("*.json")):
        with open(f, "r", encoding="utf-8") as fp:
            results.append(json.load(fp))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval nếu có)."""
    cfg = result["cfg"]
    summary = result["summary"]
    exp_id = cfg["exp_id"]

    hidden_val = cfg.get("hidden", (256, 128))
    if isinstance(hidden_val, (list, tuple)):
        hidden_str = "-".join(str(h) for h in hidden_val)
    else:
        hidden_str = str(hidden_val)

    row = {
        "exp_id": exp_id,
        "group": cfg.get("group", ""),
        "description": cfg.get("description", ""),
        "loss": cfg.get("loss", "ce").upper(),
        "optimizer": cfg.get("optimizer", ""),
        "lr": cfg.get("lr", 0.0),
        "weight_decay": cfg.get("weight_decay", 0.0),
        "batch": cfg.get("batch", 512),
        "epochs": cfg.get("epochs", 20),
        "hidden": hidden_str,
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": cfg.get("clip_norm", ""),
        "precision": cfg.get("precision", "fp32"),
        "init": cfg.get("init", "he"),
        "seed": cfg.get("seed", 1),
        "step0_loss": round(summary.get("step0_loss", 0.0), 4) if summary.get("step0_loss") is not None else "",
        "best_val_loss": round(summary.get("best_val_loss", 0.0), 4) if summary.get("best_val_loss") is not None else "",
        "best_epoch": summary.get("best_epoch", 1),
        "final_train_loss": round(summary.get("final_train_loss", 0.0), 4) if summary.get("final_train_loss") is not None else "",
        "final_val_loss": round(summary.get("final_val_loss", 0.0), 4) if summary.get("final_val_loss") is not None else "",
        "val_acc": round(summary.get("val_acc", 0.0), 4) if summary.get("val_acc") is not None else "",
        "val_macro_f1": round(summary.get("val_macro_f1", 0.0), 4) if summary.get("val_macro_f1") is not None else "",
        "time_per_epoch_s": round(summary.get("time_per_epoch_s", 0.0), 2) if summary.get("time_per_epoch_s") is not None else "",
        "peak_mem_MB": round(summary.get("peak_mem_MB", 0.0), 1) if summary.get("peak_mem_MB") is not None else 0.0,
        "diverged": "Có" if summary.get("diverged", False) else "Không",
        "eval_acc": round(eval_scores["acc"], 4) if eval_scores and "acc" in eval_scores else "",
        "eval_macro_f1": round(eval_scores["macro_f1"], 4) if eval_scores and "macro_f1" in eval_scores else "",
        "figure_file": f"figures/{exp_id}.png",
        "notes": notes or cfg.get("notes", ""),
    }
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str, summary_notes: dict | None = None) -> None:
    """Điền các dòng vào sheet 'Experiments' của mẫu và cập nhật nhận xét sheet 'Summary'."""
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]

    # Đọc header dòng 1 để map cột
    col_map = {}
    for col_idx in range(1, ws.max_column + 1):
        name = ws.cell(row=1, column=col_idx).value
        if name:
            col_map[name] = col_idx

    for r_idx, row_dict in enumerate(rows, start=2):
        for key in EXPERIMENT_COLUMNS:
            if key in col_map and key in row_dict:
                ws.cell(row=r_idx, column=col_map[key], value=row_dict[key])

        # Cập nhật công thức cho các cột ở đuôi (AD, AE, AF, AG)
        ws.cell(row=r_idx, column=30, value=f'=IF(P{r_idx}="","",P{r_idx}-LN(7))')
        ws.cell(row=r_idx, column=31, value=f'=IF(OR(T{r_idx}="",S{r_idx}=""),"",T{r_idx}-S{r_idx})')
        ws.cell(row=r_idx, column=32, value=f'=IF(OR(V{r_idx}="",Seeds!$C$8=""),"",V{r_idx}-Seeds!$C$8)')
        ws.cell(row=r_idx, column=33, value=f'=IF(OR(AF{r_idx}="",Seeds!$C$10=""),"",IF(ABS(AF{r_idx})>Seeds!$C$10,"Có","Không"))')

    # Cập nhật sheet Summary nếu có nhận xét
    if summary_notes and "Summary" in wb.sheetnames:
        ws_sum = wb["Summary"]
        for r in range(2, ws_sum.max_row + 1):
            grp = ws_sum.cell(row=r, column=1).value
            if grp in summary_notes:
                ws_sum.cell(row=r, column=8, value=summary_notes[grp])

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Đã lưu bảng thí nghiệm ra {out_path} ({len(rows)} thí nghiệm)")
